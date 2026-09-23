"""Regression tests for scouting handoffs and review workbooks.

Run with python3 -B -m unittest discover -s scripts -p 'test_dataset_scouting.py'.
Set DATASET_SCOUTING_SKILL to test the same cases against a baseline skill copy.
"""

import os
from pathlib import Path
import sys
import tempfile
import unittest

from openpyxl import load_workbook

SKILL = Path(os.environ.get(
    "DATASET_SCOUTING_SKILL",
    Path(__file__).resolve().parents[1] / "plugins/method-skills/skills/dataset-scouting",
))
sys.path.insert(0, str(SKILL / "scripts"))

from datasheet_lib.markdown import candidate_scaffold, parse_candidate
from datasheet_lib.model import SUMMARY_KEYS
from datasheet_lib.validate import validate_datasheet
from datasheet_lib.workbook import build_workbook


class DatasheetTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        (self.root / "candidates").mkdir()
        (self.root / "criteria.md").write_text(
            "## Fields of interest\n\n"
            "| Field | Priority | Threshold |\n|---|---|---|\n"
            "| cell population | must-have | 100% |\n\n"
            "## Technical details\n\n- alignment\n",
            encoding="utf-8",
        )
        summary = {
            "Identifiers": "ALPHA",
            "Publication": "Fixture study [S1]",
            "Modality": "scRNA-seq [S1]",
            "Sample count": "2 samples [S1]",
            "Download size": "1 MB total compressed [S1]",
            "Metadata route": "https://archive.example/ALPHA/samples.tsv [S1]",
            "Metadata access": "direct",
            "Metadata evidence": "contents inspected 2026-09-23 [S1]",
            "Data source": "GEO / S3 [S1]",
            "Download method": "AWS CLI for the listed S3 objects [S1]",
            "Raw-data route": "s3://scouting-fixture/ALPHA/counts.mtx.gz [S1]",
            "Raw-data access": "direct",
            "Raw-data evidence": "contents inspected 2026-09-23 [S1]",
            "Access and contacts": "Public objects; no application needed [S1]",
            "Tier": "1",
            "Recommendation": "accept — counts and both donor arms are available [S1]",
        }
        text = candidate_scaffold("ALPHA")
        # Use each version's scaffold so baseline failures measure behavior, not new keys.
        for key in SUMMARY_KEYS:
            text = text.replace(f"- **{key}:**", f"- **{key}:** {summary[key]}")
        for name in ("Metadata", "Required fields", "Raw data"):
            text = text.replace(f"| {name} |  |  |  |", f"| {name} | pass | fixture inspected | [S1] |")
        text = text.replace(
            "|  |  |  |  |  |  |  |  |",
            '| cell population | must-have | present | 2/2 | export | sample | [S1] | CD4 T cells; source: "CD4 memory T cells", memory-selected |',
        )
        text = text.replace("|  |  |  |  |", "| alignment | Cell Ranger 7.1 / GRCh38 | [S1] | publication |", 1)
        text = text.replace(
            "|  |  |  |",
            "| publication | full fixture read | [S1] |\n| supplements | none in this fixture | [S1] |", 1,
        )
        text = text.replace("|  |  |  |  |", "| S1 | supplied fixture | https://archive.example/ALPHA/record | 2026-09-23 |", 1)
        text = text.replace("- **Notes:**", "- **Notes:** Keep my wording, including café and  two spaces.")
        self.path = self.root / "candidates/ALPHA.md"
        self.path.write_text(text, encoding="utf-8")

    def replace(self, old, new):
        text = self.path.read_text(encoding="utf-8")
        self.assertIn(old, text)
        self.path.write_text(text.replace(old, new), encoding="utf-8")

    def errors(self, baseline=None):
        _, _, issues = validate_datasheet(str(self.root), baseline)
        return [issue.message for issue in issues if issue.severity == "ERROR"]

    def build(self):
        candidates, criteria, issues = validate_datasheet(str(self.root))
        self.assertEqual([i.message for i in issues if i.severity == "ERROR"], [])
        out = build_workbook(str(self.root), candidates, criteria)
        book = load_workbook(out)
        self.addCleanup(book.close)
        return book

    def main_values(self, book):
        return dict(zip(next(book["Datasheet"].values), list(book["Datasheet"].values)[1]))

    def test_valid_candidate(self):
        self.assertEqual(self.errors(), [])

    def test_prose_is_not_a_passing_raw_route(self):
        self.replace("s3://scouting-fixture/ALPHA/counts.mtx.gz", "GEO FTP; SRA per sample")
        self.assertTrue(self.errors())

    def test_prose_is_not_a_passing_metadata_route(self):
        self.replace("https://archive.example/ALPHA/samples.tsv", "See GEO characteristics")
        self.assertTrue(self.errors())

    def test_study_accession_is_not_a_download_run(self):
        self.replace("s3://scouting-fixture/ALPHA/counts.mtx.gz", "SRP123456")
        self.assertTrue(self.errors())

    def test_mixed_protocol_targets_survive_export(self):
        targets = [
            "s3://scouting-fixture/ALPHA/counts.mtx.gz",
            "ftp://archive.example/ALPHA/features.tsv.gz",
            "https://archive.example/ALPHA/barcodes(v1).tsv.gz?download=1&token=abc%2Fxyz",
        ]
        self.replace("s3://scouting-fixture/ALPHA/counts.mtx.gz [S1]", "\n".join(t + " [S1]" for t in targets))
        if "Download method" in SUMMARY_KEYS:
            self.replace("AWS CLI for the listed S3 objects", "AWS CLI for counts; FTP/HTTPS for features and barcodes")
        self.assertEqual(self.main_values(self.build())["Raw-data route"], "\n".join(targets))

    def test_run_identifiers_survive_export(self):
        self.replace("s3://scouting-fixture/ALPHA/counts.mtx.gz [S1]", "SRR123456 [S1]\nERR123457 [S1]")
        if "Download method" in SUMMARY_KEYS:
            self.replace("AWS CLI for the listed S3 objects", "SRA Toolkit: prefetch then fasterq-dump for each run")
        self.assertEqual(self.main_values(self.build())["Raw-data route"], "SRR123456\nERR123457")

    def test_single_https_target_remains_clickable(self):
        target = "https://archive.example/ALPHA/counts(v1).h5ad?download=1&token=abc%2Fxyz"
        self.replace("s3://scouting-fixture/ALPHA/counts.mtx.gz", target)
        book = self.build()
        column = list(next(book["Datasheet"].values)).index("Raw-data route") + 1
        cell = book["Datasheet"].cell(2, column)
        self.assertEqual(cell.value, target)
        self.assertEqual(cell.hyperlink.target, target)

    def test_markdown_link_destinations_remain_supported(self):
        target = "https://archive.example/ALPHA/counts.h5ad?download=1"
        for route in (f"[download]({target})", f"[download](<{target}>)", f"`{target}`", f"<{target}>"):
            with self.subTest(route=route):
                self.replace("s3://scouting-fixture/ALPHA/counts.mtx.gz", route)
                self.assertEqual(self.main_values(self.build())["Raw-data route"], target)
                self.replace(route, "s3://scouting-fixture/ALPHA/counts.mtx.gz")

    def test_gen3_file_identifiers_survive_export(self):
        guid = "00149bcf-e057-4ecc-b22d-53648ae0b35f"
        for target in (guid, f"dg.4DFC/{guid}"):
            with self.subTest(target=target):
                self.replace("s3://scouting-fixture/ALPHA/counts.mtx.gz", target)
                if "Download method" in SUMMARY_KEYS:
                    self.replace("AWS CLI for the listed S3 objects", "gen3-client download-single --guid with the listed file GUID")
                self.assertEqual(self.main_values(self.build())["Raw-data route"], target)
                self.replace(target, "s3://scouting-fixture/ALPHA/counts.mtx.gz")
                if "Download method" in SUMMARY_KEYS:
                    self.replace("gen3-client download-single --guid with the listed file GUID", "AWS CLI for the listed S3 objects")

    def test_local_manifest_must_exist(self):
        self.replace("s3://scouting-fixture/ALPHA/counts.mtx.gz", "./manifest.tsv")
        self.assertTrue(self.errors())
        (self.root / "manifest.tsv").write_text("url\ns3://scouting-fixture/ALPHA/counts.mtx.gz\n")
        self.assertEqual(self.errors(), [])

    def test_unknown_method_cannot_support_pass(self):
        if "Download method" not in SUMMARY_KEYS:
            self.skipTest("baseline has no download-method contract")
        self.replace("AWS CLI for the listed S3 objects", "unknown — not resolved")
        self.assertTrue(self.errors())

    def test_restricted_data_cannot_pass(self):
        self.replace("- **Raw-data access:** direct", "- **Raw-data access:** on request")
        self.assertTrue(self.errors())

    def test_restricted_candidate_can_remain_unknown(self):
        self.replace("- **Raw-data access:** direct", "- **Raw-data access:** on request")
        self.replace("- **Raw-data evidence:** contents inspected 2026-09-23", "- **Raw-data evidence:** claimed")
        self.replace("s3://scouting-fixture/ALPHA/counts.mtx.gz", "https://archive.example/ALPHA/apply")
        self.replace("| Raw data | pass |", "| Raw data | unknown |")
        self.replace("accept — counts and both donor arms are available", "unknown — application needed for the data")
        self.assertEqual(self.errors(), [])
        self.assertTrue(self.main_values(self.build())["Recommendation"].startswith("unknown"))

    def test_existing_evidence_and_coverage_gates(self):
        self.replace("| 2/2 |", "| 1/2 |")
        self.assertTrue(self.errors())
        self.replace("| 1/2 |", "| 2/2 |")
        self.replace("- **Raw-data evidence:** contents inspected 2026-09-23", "- **Raw-data evidence:** route verified 2026-09-23")
        self.assertTrue(self.errors())
        self.replace("| Raw data | pass |", "| Raw data | unknown |")
        self.assertTrue(self.errors())  # An unresolved check must still prevent accept.

    def test_report_contains_values_reason_and_source_method(self):
        book = self.build()
        self.assertEqual(book.sheetnames, ["Datasheet", "Fields", "Technical", "Totals"])
        values = self.main_values(book)
        self.assertIn("both donor arms", values["Recommendation"])
        self.assertEqual(values["Data source"], "GEO / S3")
        self.assertIn("AWS CLI", values["Download method"])
        field = book["Fields"].cell(2, 3).value
        self.assertIn("CD4 T cells", field)
        self.assertIn("memory-selected", field)
        self.assertIn("2/2", field)
        self.assertIn("sample", field)
        self.assertIn("Cell Ranger 7.1", book["Technical"].cell(2, 2).value)

    def test_build_preserves_review_and_baseline_detects_change(self):
        before = self.path.read_bytes()
        baseline = self.root / "baseline"
        (baseline / "candidates").mkdir(parents=True)
        (baseline / "candidates/ALPHA.md").write_bytes(before)
        self.build()
        self.assertEqual(self.path.read_bytes(), before)
        self.assertEqual(self.errors(str(baseline)), [])
        self.replace("Keep my wording", "Changed wording")
        self.assertTrue(self.errors(str(baseline)))

    def test_legacy_record_parses_without_migration(self):
        if "Data source" not in SUMMARY_KEYS:
            self.skipTest("baseline already uses legacy fields")
        text = self.path.read_text(encoding="utf-8")
        text = "\n".join(line for line in text.splitlines() if not line.startswith(("- **Data source:", "- **Download method:")))
        self.path.write_text(text, encoding="utf-8")
        review = parse_candidate(self.path).review_text
        self.assertTrue(self.errors())
        self.assertEqual(parse_candidate(self.path).review_text, review)
        self.assertEqual(self.path.read_text(encoding="utf-8"), text)


if __name__ == "__main__":
    unittest.main()
